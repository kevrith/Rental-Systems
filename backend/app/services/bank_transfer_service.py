"""Bank transfer as a payment method, and bank statement reconciliation
(Sprint 23, US-101).

Recording a bank transfer a caretaker was told about is nothing new — it is
the same `payment_service.record_cash_payment` cash and cheques already use,
just with `PaymentMethod.BANK_TRANSFER` and the transfer's own reference
stored on `Payment.bank_reference` rather than the M-Pesa-only
`mpesa_receipt` column.

What is new is going the other direction: a landlord who banks rent directly
uploads their monthly statement, and every line either names a tenancy (its
reference code appears somewhere in the bank's own narrative) or it does not.
Matching only ever *proposes* a tenancy — recording the payment is still an
explicit choice per row, the same two-step preview-then-commit shape
`import_service` already uses for the tenant bulk import, and for the same
reason: nothing is written until a person has seen exactly what will be.
"""

import csv
import io
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from fastapi import HTTPException, Request, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import OrgContext, assert_in_org
from app.models.billing import PaymentMethod
from app.models.integrations import BankStatementEntry, BankStatementUpload
from app.models.organization import Organization
from app.models.tenant import Tenancy
from app.services import payment_service

MAX_ROWS = 1000

_DATE_ALIASES = {"date", "transaction date", "value date", "posting date"}
_DESCRIPTION_ALIASES = {"description", "narrative", "details", "particulars", "narration"}
_AMOUNT_ALIASES = {"amount", "credit", "credit amount", "deposit"}

# A RentFlow tenancy reference embedded anywhere in a bank's own narrative,
# e.g. "RENT TCY-7K2M9Q JOHN DOE" (see `reference_service`'s alphabet).
_REFERENCE_PATTERN = re.compile(r"\bTCY-[A-Z2-9]{6}\b")


def instructions(organization: Organization) -> dict:
    """What a tenant sees under "Pay by bank transfer" (US-101)."""
    return {
        "configured": bool(organization.bank_account_number),
        "bank_name": organization.bank_name,
        "account_name": organization.bank_account_name,
        "account_number": organization.bank_account_number,
        "branch": organization.bank_branch,
    }


# ------------------------------------------------------------------- parsing


def _header_index(header: list[str], aliases: set[str]) -> int | None:
    for index, name in enumerate(header):
        if name in aliases:
            return index
    return None


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def _parse_amount(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        return Decimal(str(value).replace(",", "").replace("KES", "").strip())
    except (InvalidOperation, ValueError):
        return None


def _rows_from_csv(data: bytes) -> tuple[list[str], list[list[Any]]]:
    text = data.decode("utf-8-sig", errors="ignore")
    rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    return (rows[0], rows[1:]) if rows else ([], [])


def _rows_from_excel(data: bytes) -> tuple[list[Any], list[list[Any]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    rows = [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    return (rows[0], rows[1:]) if rows else ([], [])


def parse_statement(data: bytes, filename: str) -> tuple[list[dict], list[dict]]:
    """Read a CSV or Excel bank statement into (rows, errors). Nothing is written.

    Only credit lines (money coming in) matter for rent reconciliation, so a
    debit or zero-amount row is silently dropped rather than reported as an
    error.
    """
    try:
        header_row, body = (
            _rows_from_csv(data) if filename.lower().endswith(".csv") else _rows_from_excel(data)
        )
    except Exception as exc:  # noqa: BLE001 — any parse failure is the same 400 to the caller
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That file could not be read as a CSV or Excel statement",
        ) from exc

    if not header_row:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The statement is empty")

    header = [str(cell or "").strip().lower() for cell in header_row]
    date_index = _header_index(header, _DATE_ALIASES)
    description_index = _header_index(header, _DESCRIPTION_ALIASES)
    amount_index = _header_index(header, _AMOUNT_ALIASES)
    missing = [
        label
        for label, index in (
            ("a date", date_index),
            ("a description", description_index),
            ("an amount/credit", amount_index),
        )
        if index is None
    ]
    if missing:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not find {', '.join(missing)} column in the statement",
        )
    # `missing` above already proved these are set — restated so mypy narrows
    # them from `int | None` to `int` for the indexing below.
    assert date_index is not None
    assert description_index is not None
    assert amount_index is not None

    rows: list[dict] = []
    errors: list[dict] = []
    for number, row in enumerate(body, start=2):
        if len(rows) + len(errors) >= MAX_ROWS:
            errors.append({"row": number, "reason": f"Statements are limited to {MAX_ROWS} rows"})
            break

        entry_date = _parse_date(row[date_index] if date_index < len(row) else None)
        description = str(row[description_index] or "").strip() if description_index < len(row) else ""
        amount = _parse_amount(row[amount_index] if amount_index < len(row) else None)

        if entry_date is None or not description or amount is None:
            errors.append(
                {"row": number, "reason": "Could not read a date, description and amount from this row"}
            )
            continue
        if amount <= 0:
            continue

        match = _REFERENCE_PATTERN.search(description.upper())
        rows.append(
            {
                "row": number,
                "entry_date": entry_date,
                "description": description,
                "amount": amount,
                "reference_guess": match.group(0) if match else None,
            }
        )

    return rows, errors


async def match_rows(db: AsyncSession, context: OrgContext, rows: list[dict]) -> list[dict]:
    """Attach a tenancy suggestion to every row whose narrative named one."""
    codes = {row["reference_guess"] for row in rows if row.get("reference_guess")}
    tenancies: dict[str, Tenancy] = {}
    if codes:
        found = await db.scalars(
            select(Tenancy).where(
                Tenancy.organization_id == context.organization_id, Tenancy.reference_code.in_(codes)
            )
        )
        tenancies = {tenancy.reference_code: tenancy for tenancy in found}

    matched = []
    for row in rows:
        tenancy = tenancies.get(row.get("reference_guess") or "")
        matched.append({**row, "matched_tenancy_id": tenancy.id if tenancy else None})
    return matched


# ---------------------------------------------------------------------- commit


async def commit_statement(
    db: AsyncSession,
    context: OrgContext,
    *,
    rows: list[Any],
    request: Request | None = None,
) -> tuple[BankStatementUpload, list[dict]]:
    """Persist every reviewed row, recording a payment for whichever the
    operator chose to confirm. A payment failure on one row never rolls back
    the rows already recorded — the same independence `import_service` gives
    the bulk tenant import.
    """
    upload = BankStatementUpload(
        organization_id=context.organization_id,
        uploaded_by_id=context.user.id,
        row_count=len(rows),
    )
    db.add(upload)
    await db.flush()

    matched_count = 0
    unmatched_count = 0
    failures: list[dict] = []

    for row in rows:
        tenancy_id = row.tenancy_id
        if tenancy_id is not None:
            assert_in_org(await db.get(Tenancy, tenancy_id), context, label="tenancy")

        entry = BankStatementEntry(
            organization_id=context.organization_id,
            upload_id=upload.id,
            entry_date=row.entry_date,
            description=row.description,
            amount=row.amount,
            matched_tenancy_id=tenancy_id,
            is_matched=tenancy_id is not None,
        )
        db.add(entry)
        await db.flush()

        if tenancy_id is None:
            unmatched_count += 1
            continue
        matched_count += 1

        if not row.record_payment:
            continue
        try:
            payment = await payment_service.record_cash_payment(
                db,
                context,
                tenancy_id=tenancy_id,
                amount=row.amount,
                payment_date=row.entry_date,
                method=PaymentMethod.BANK_TRANSFER,
                reference=row.description[:64],
                notes="Recorded from an uploaded bank statement",
                request=request,
            )
            entry.matched_payment_id = payment.id
        except HTTPException as exc:
            failures.append({"row": row.row, "reason": str(exc.detail)})

    upload.matched_count = matched_count
    upload.unmatched_count = unmatched_count
    await db.commit()
    await db.refresh(upload)
    return upload, failures


async def list_uploads(db: AsyncSession, context: OrgContext, limit: int = 50) -> list[BankStatementUpload]:
    return list(
        await db.scalars(
            select(BankStatementUpload)
            .where(BankStatementUpload.organization_id == context.organization_id)
            .order_by(BankStatementUpload.created_at.desc())
            .limit(limit)
        )
    )


async def get_upload(db: AsyncSession, context: OrgContext, upload_id: uuid.UUID) -> BankStatementUpload:
    upload = assert_in_org(
        await db.scalar(
            select(BankStatementUpload)
            .options(selectinload(BankStatementUpload.entries))
            .where(BankStatementUpload.id == upload_id)
        ),
        context,
        label="statement upload",
    )
    return upload
