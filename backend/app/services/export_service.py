"""Data export — the promise that you are never locked in (US-077).

Every export is built from a single row-shaping function per dataset, so the CSV
and the Excel workbook can never disagree about what a "payment" export contains.
Large exports are generated in the background and the requester is told when the
file is ready, rather than holding a request open for a minute.
"""

import csv
import io
import logging
import uuid
from collections.abc import Awaitable, Callable
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import OrgContext, accessible_property_ids
from app.models.billing import Invoice, Payment
from app.models.file import FileCategory
from app.models.notification import NotificationChannel, NotificationType
from app.models.property import Property, Unit
from app.models.tenant import Tenancy, Tenant
from app.models.vacancy import DataExport, ExportFormat, ExportKind
from app.services import audit_service, file_service, notification_service

logger = logging.getLogger(__name__)

ZERO = Decimal("0.00")

# Above this, the export is built in the background and collected later rather
# than streamed back on the request.
BACKGROUND_THRESHOLD = 2000


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    if value is None:
        return ""
    if hasattr(value, "value"):  # enum
        return value.value
    return value


async def _tenant_rows(db: AsyncSession, context: OrgContext, **_) -> list[dict]:
    rows = await db.scalars(
        select(Tenant).where(Tenant.organization_id == context.organization_id).order_by(Tenant.full_name)
    )
    return [
        {
            "Reference": tenant.reference_code,
            "Full name": tenant.full_name,
            "Phone": tenant.phone_number,
            "Email": tenant.email,
            "National ID": tenant.national_id,
            "Employer": tenant.employer_name,
            "Occupation": tenant.occupation,
            "Monthly income": tenant.monthly_income,
            "Emergency contact": tenant.emergency_contact_name,
            "Emergency phone": tenant.emergency_contact_phone,
            "Archived": tenant.is_archived,
            "Added on": tenant.created_at,
        }
        for tenant in rows
    ]


async def _tenancy_rows(db: AsyncSession, context: OrgContext, **_) -> list[dict]:
    rows = await db.execute(
        select(Tenancy, Tenant, Unit, Property)
        .join(Tenant, Tenant.id == Tenancy.tenant_id)
        .join(Unit, Unit.id == Tenancy.unit_id)
        .join(Property, Property.id == Unit.property_id)
        .where(Tenancy.organization_id == context.organization_id)
        .order_by(Tenancy.start_date.desc())
    )
    return [
        {
            "Reference": tenancy.reference_code,
            "Tenant": tenant.full_name,
            "Phone": tenant.phone_number,
            "Property": property_record.name,
            "Unit": unit.unit_number,
            "Status": tenancy.status,
            "Start": tenancy.start_date,
            "End": tenancy.end_date,
            "Open ended": tenancy.is_open_ended,
            "Monthly rent": tenancy.monthly_rent,
            "Deposit": tenancy.deposit_amount,
            "Billing day": tenancy.billing_day,
            "Notice period (days)": tenancy.notice_period_days,
        }
        for tenancy, tenant, unit, property_record in rows.all()
    ]


async def _payment_rows(
    db: AsyncSession, context: OrgContext, *, date_from: date | None = None, date_to: date | None = None
) -> list[dict]:
    query = (
        select(Payment, Tenant, Unit, Property)
        .join(Tenancy, Tenancy.id == Payment.tenancy_id)
        .join(Tenant, Tenant.id == Tenancy.tenant_id)
        .join(Unit, Unit.id == Tenancy.unit_id)
        .join(Property, Property.id == Unit.property_id)
        .where(Payment.organization_id == context.organization_id)
    )
    if date_from:
        query = query.where(Payment.payment_date >= date_from)
    if date_to:
        query = query.where(Payment.payment_date <= date_to)

    rows = await db.execute(query.order_by(Payment.payment_date.desc()))
    return [
        {
            "Reference": payment.reference_code,
            "Date": payment.payment_date,
            "Tenant": tenant.full_name,
            "Property": property_record.name,
            "Unit": unit.unit_number,
            "Amount": payment.amount,
            "Method": payment.method,
            "Status": payment.status,
            "M-Pesa code": payment.mpesa_receipt_number,
            "Recorded on": payment.created_at,
        }
        for payment, tenant, unit, property_record in rows.all()
    ]


async def _property_rows(db: AsyncSession, context: OrgContext, **_) -> list[dict]:
    query = select(Property).where(Property.organization_id == context.organization_id)
    allowed = await accessible_property_ids(db, context)
    if allowed is not None:
        query = query.where(Property.id.in_(allowed))

    rows = await db.scalars(query.order_by(Property.name))
    return [
        {
            "Reference": record.reference_code,
            "Name": record.name,
            "Type": record.property_type,
            "Address": record.address,
            "County": record.county,
            "Sub-county": record.sub_county,
            "Water rate": record.water_rate_per_unit,
            "Electricity rate": record.electricity_rate_per_unit,
            "Grace period (days)": record.grace_period_days,
            "Maintenance budget": record.maintenance_budget_monthly,
            "Archived": record.is_archived,
        }
        for record in rows
    ]


async def _unit_rows(db: AsyncSession, context: OrgContext, **_) -> list[dict]:
    rows = await db.execute(
        select(Unit, Property)
        .join(Property, Property.id == Unit.property_id)
        .where(Unit.organization_id == context.organization_id)
        .order_by(Property.name, Unit.unit_number)
    )
    return [
        {
            "Reference": unit.reference_code,
            "Property": property_record.name,
            "Unit": unit.unit_number,
            "Type": unit.unit_type,
            "Use class": unit.use_class,
            "Floor": unit.floor,
            "Size (sqm)": unit.size_sqm,
            "Bedrooms": unit.bedrooms,
            "Bathrooms": unit.bathrooms,
            "Parking bays": unit.car_bays,
            "Monthly rent": unit.monthly_rent,
            "Deposit": unit.deposit_amount,
            "Status": unit.status,
        }
        for unit, property_record in rows.all()
    ]


async def _invoice_rows(
    db: AsyncSession, context: OrgContext, *, date_from: date | None = None, date_to: date | None = None
) -> list[dict]:
    query = (
        select(Invoice, Tenant, Unit, Property)
        .join(Tenancy, Tenancy.id == Invoice.tenancy_id)
        .join(Tenant, Tenant.id == Tenancy.tenant_id)
        .join(Unit, Unit.id == Tenancy.unit_id)
        .join(Property, Property.id == Unit.property_id)
        .where(Invoice.organization_id == context.organization_id)
    )
    if date_from:
        query = query.where(Invoice.issue_date >= date_from)
    if date_to:
        query = query.where(Invoice.issue_date <= date_to)

    rows = await db.execute(query.order_by(Invoice.issue_date.desc()))
    return [
        {
            "Reference": invoice.reference_code,
            "Issued": invoice.issue_date,
            "Due": invoice.due_date,
            "Tenant": tenant.full_name,
            "Property": property_record.name,
            "Unit": unit.unit_number,
            "Period start": invoice.period_start,
            "Period end": invoice.period_end,
            "Total": invoice.total,
            "Paid": invoice.amount_paid,
            "Balance": Decimal(invoice.total) - Decimal(invoice.amount_paid),
            "Status": invoice.status,
        }
        for invoice, tenant, unit, property_record in rows.all()
    ]


# Every dataset is one row-shaping function, so CSV and Excel can never disagree
# about what an export of that dataset contains.
RowBuilder = Callable[..., Awaitable[list[dict]]]

BUILDERS: dict[ExportKind, RowBuilder] = {
    ExportKind.TENANTS: _tenant_rows,
    ExportKind.TENANCIES: _tenancy_rows,
    ExportKind.PAYMENTS: _payment_rows,
    ExportKind.PROPERTIES: _property_rows,
    ExportKind.UNITS: _unit_rows,
    ExportKind.INVOICES: _invoice_rows,
}


def to_csv(rows: list[dict]) -> bytes:
    if not rows:
        return b""
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
    writer.writeheader()
    for row in rows:
        writer.writerow({key: _cell(value) for key, value in row.items()})
    return buffer.getvalue().encode("utf-8-sig")


def to_excel(rows: list[dict], sheet_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name[:31]

    if rows:
        headers = list(rows[0].keys())
        fill = PatternFill("solid", fgColor="1F3A93")
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 4)

        for row_index, row in enumerate(rows, start=2):
            for column_index, header in enumerate(headers, start=1):
                sheet.cell(row=row_index, column=column_index, value=_cell(row[header]))

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def build(
    db: AsyncSession,
    context: OrgContext,
    kind: ExportKind,
    export_format: ExportFormat,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[bytes, str, int]:
    """Return (bytes, filename, row count) for an export the caller can stream."""
    rows = await BUILDERS[kind](db, context, date_from=date_from, date_to=date_to)
    stamp = date.today().isoformat()

    if export_format == ExportFormat.CSV:
        return to_csv(rows), f"rentflow-{kind.value}-{stamp}.csv", len(rows)
    return (
        to_excel(rows, kind.value.title()),
        f"rentflow-{kind.value}-{stamp}.xlsx",
        len(rows),
    )


async def record_export(
    db: AsyncSession,
    context: OrgContext,
    kind: ExportKind,
    export_format: ExportFormat,
    row_count: int,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
    file_id: uuid.UUID | None = None,
    is_scheduled: bool = False,
    record_id: uuid.UUID | None = None,
) -> DataExport:
    """Keep a log of who exported what — this is the whole tenant book leaving."""
    record = DataExport(
        id=record_id or uuid.uuid4(),
        organization_id=context.organization_id,
        kind=kind,
        export_format=export_format,
        date_from=date_from,
        date_to=date_to,
        row_count=row_count,
        file_id=file_id,
        is_scheduled=is_scheduled,
        requested_by_id=context.user.id if context.user else None,
    )
    db.add(record)

    audit_service.record(
        db,
        organization_id=context.organization_id,
        action="data.exported",
        entity_type="data_export",
        entity_id=record.id,
        actor=context.user,
        summary=f"Exported {row_count} {kind.value} row(s) as {export_format.value}",
    )
    await db.commit()
    await db.refresh(record)
    return record


async def build_and_store(
    db: AsyncSession,
    organization_id: uuid.UUID,
    kind: ExportKind,
    export_format: ExportFormat,
    *,
    context: OrgContext,
    notify_user_id: uuid.UUID | None = None,
) -> DataExport:
    """The background path: build the file, store it, and say when it is ready."""
    data, filename, row_count = await build(db, context, kind, export_format)
    # The file is filed against the export row it belongs to, so a stored export
    # can always be traced back to who asked for it and when.
    record_id = uuid.uuid4()
    stored = await file_service.register_generated(
        db,
        organization_id,
        data=data,
        filename=filename,
        category=FileCategory.OTHER,
        entity_type="data_export",
        entity_id=record_id,
    )
    record = await record_export(
        db, context, kind, export_format, row_count, file_id=stored.id, record_id=record_id
    )

    if notify_user_id is not None:
        from app.models.user import User

        user = await db.get(User, notify_user_id)
        if user is not None:
            await notification_service.send(
                db,
                recipient=notification_service.Recipient.for_user(user),
                notification_type=NotificationType.EXPORT_READY,
                title="Your export is ready",
                body=f"{row_count} {kind.value} row(s) exported as {filename}.",
                channels=[NotificationChannel.IN_APP, NotificationChannel.EMAIL],
                link_path="/settings/exports",
                entity_type="data_export",
                entity_id=record.id,
            )
            await db.commit()

    return record


async def run_scheduled_exports(db: AsyncSession) -> int:
    """The monthly backup every owner gets whether they ask or not (US-077).

    One tenancies workbook per organisation, filed in their documents and
    announced in-app. It exists so that a customer who walks away still has
    their data, without having to remember to take it.
    """
    from app.models.organization import Organization
    from app.models.user import User, UserRole

    organizations = list(await db.scalars(select(Organization).where(Organization.is_active.is_(True))))

    built = 0
    for organization in organizations:
        owner = await db.scalar(
            select(User)
            .where(
                User.organization_id == organization.id,
                User.role.in_([UserRole.OWNER, UserRole.AGENCY_ADMIN]),
                User.is_active.is_(True),
            )
            .order_by(User.created_at)
        )
        if owner is None:
            continue

        context = OrgContext(user=owner, organization=organization)
        try:
            await build_and_store(
                db,
                organization.id,
                ExportKind.TENANCIES,
                ExportFormat.EXCEL,
                context=context,
                notify_user_id=owner.id,
            )
            built += 1
        except Exception:  # noqa: BLE001 — one organisation must not stop the rest
            logger.exception("Scheduled export failed for organisation %s", organization.id)
            await db.rollback()

    return built
